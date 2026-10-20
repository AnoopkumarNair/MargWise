"""Build docs/golden/golden-cases.xlsx: money/assessment golden cases (03 §3.13).

Every expected value is a spreadsheet FORMULA over visible input cells, so the
workbook is an implementation independent of the TypeScript engine. Python here
only lays out inputs and formulas; it computes nothing. After building, run
LibreOffice recalculation, then export_golden.py to produce the JSON the engine
tests assert against (within ±₹1 or ±0.01%).

Tax rules on the Rules sheet are the values written in 03 §3.8, used here as
TEST INPUTS. In the product they come from the reviewed content release.

Usage: python3 -I scripts/golden/build_golden.py docs/golden/golden-cases.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BLUE = Font(name="Arial", color="0000FF")
BLACK = Font(name="Arial", color="000000")
BOLD = Font(name="Arial", bold=True)
YELLOW = PatternFill("solid", fgColor="FFFF00")
INR = "#,##0.00"

wb = Workbook()
readme = wb.active
readme.title = "README"
for i, line in enumerate(
    [
        "Margwise golden cases — STATUS: DRAFT, awaiting Anu's approval (then a CA's review).",
        "Blue = input. Black = formula. Yellow = expected value asserted by the engine tests.",
        "Expected values are spreadsheet formulas only; they never come from app code.",
        "Rules sheet: tax values from spec 03 §3.8 used as test inputs (production values come from the content release).",
        "Golden values change only when this workbook changes (CI check). Never edit the JSON by hand.",
        "Tolerance: ±₹1 or ±0.01%, whichever is larger.",
        "Open question for Anu/CA: new-regime marginal relief above ₹12L (case TAX-N-REL) is not in 03 §3.8.",
    ],
    start=1,
):
    readme.cell(row=i, column=1, value=line).font = BOLD if i == 1 else BLACK


def inp(ws, ref, value, note=None):
    c = ws[ref]
    c.value = value
    c.font = BLUE
    if note:
        c.comment = Comment(note, "golden")
    return c


def fx(ws, ref, formula, expected=False):
    c = ws[ref]
    c.value = formula
    c.font = BLACK
    if expected:
        c.fill = YELLOW
    return c


# ----------------------------------------------------------------- Rules
rules = wb.create_sheet("Rules")
rules["A1"] = "Regime"
rules["B1"] = "From (₹)"
rules["C1"] = "To (₹, exclusive)"
rules["D1"] = "Rate"
new_slabs = [(0, 400000, 0), (400000, 800000, 0.05), (800000, 1200000, 0.10), (1200000, 1600000, 0.15),
             (1600000, 2000000, 0.20), (2000000, 2400000, 0.25), (2400000, 1e15, 0.30)]
old_slabs = [(0, 250000, 0), (250000, 500000, 0.05), (500000, 1000000, 0.20), (1000000, 1e15, 0.30)]
r = 2
for f, t, rate in new_slabs:
    rules.cell(row=r, column=1, value="new")
    inp(rules, f"B{r}", f); inp(rules, f"C{r}", t); inp(rules, f"D{r}", rate)
    r += 1
for f, t, rate in old_slabs:
    rules.cell(row=r, column=1, value="old")
    inp(rules, f"B{r}", f); inp(rules, f"C{r}", t); inp(rules, f"D{r}", rate)
    r += 1
NEW = ("Rules!$B$2:$B$8", "Rules!$C$2:$C$8", "Rules!$D$2:$D$8")
OLD = ("Rules!$B$9:$B$12", "Rules!$C$9:$C$12", "Rules!$D$9:$D$12")
rules["F1"] = "Parameter"; rules["G1"] = "new"; rules["H1"] = "old"
params = [("standard_deduction", 75000, 50000), ("rebate_max_taxable", 1200000, 500000), ("rebate_max", 60000, 12500), ("cess", 0.04, 0.04)]
for i, (name, n, o) in enumerate(params, start=2):
    rules[f"F{i}"] = name
    inp(rules, f"G{i}", n, "03 §3.8 (test input)")
    inp(rules, f"H{i}", o, "03 §3.8 (test input)")
SD = {"new": "Rules!$G$2", "old": "Rules!$H$2"}
RB_LIM = {"new": "Rules!$G$3", "old": "Rules!$H$3"}
RB_MAX = {"new": "Rules!$G$4", "old": "Rules!$H$4"}
CESS = {"new": "Rules!$G$5", "old": "Rules!$H$5"}


def slab_tax(taxable, regime):
    f, t, rate = NEW if regime == "new" else OLD
    return f"SUMPRODUCT(({taxable}>{f})*((({taxable}<{t})*{taxable}+({taxable}>={t})*{t})-{f})*{rate})"


def income_tax(gross, regime, relief="0"):
    """Total tax incl. cess as one formula string (mirrors 03 §3.8)."""
    taxable = f"MAX(0,{gross}-{SD[regime]})"
    st = slab_tax(taxable, regime)
    rebate = f"IF({taxable}<={RB_LIM[regime]},MIN({st},{RB_MAX[regime]}),0)"
    rel = f"IF(AND({relief}=1,{taxable}>{RB_LIM[regime]}),MAX(0,{st}-{rebate}-({taxable}-{RB_LIM[regime]})),0)"
    return f"({st}-{rebate}-{rel})*(1+{CESS[regime]})"


def marginal_rate(gross, regime):
    f, t, rate = NEW if regime == "new" else OLD
    taxable = f"MAX(0,{gross}-{SD[regime]})"
    return f"SUMPRODUCT(({taxable}>={f})*({taxable}<{t})*{rate})*(1+{CESS[regime]})"


# ----------------------------------------------------------------- Tax
tax = wb.create_sheet("Tax")
for col, h in enumerate(["id", "regime", "gross", "marginal_relief", "expected_total"], start=1):
    tax.cell(row=1, column=col, value=h).font = BOLD
tax_cases = [
    ("TAX-N-4L", "new", 475000, 0), ("TAX-N-8L", "new", 875000, 0), ("TAX-N-12L", "new", 1275000, 0),
    ("TAX-N-12L+1", "new", 1275001, 0), ("TAX-N-REL", "new", 1300000, 1), ("TAX-N-24L", "new", 2475000, 0),
    ("TAX-N-30L", "new", 3075000, 0), ("TAX-O-2.5L", "old", 300000, 0), ("TAX-O-5L", "old", 550000, 0),
    ("TAX-O-5L+1", "old", 550001, 0), ("TAX-O-10L", "old", 1050000, 0), ("TAX-O-20L", "old", 2050000, 0),
]
for i, (cid, regime, gross, relief) in enumerate(tax_cases, start=2):
    tax[f"A{i}"] = cid
    inp(tax, f"B{i}", regime)
    inp(tax, f"C{i}", gross)
    inp(tax, f"D{i}", relief)
    fx(tax, f"E{i}", "=" + income_tax(f"C{i}", regime, f"D{i}"), expected=True).number_format = INR

# ----------------------------------------------------------------- EMI
emi = wb.create_sheet("EMI")
for col, h in enumerate(["id", "principal", "annual_rate", "months", "expected_emi"], start=1):
    emi.cell(row=1, column=col, value=h).font = BOLD
for i, (cid, p, rate, n) in enumerate(
    [("EMI-10L-10%-10y", 1000000, 0.10, 120), ("EMI-7.5L-10.15%-7y", 750000, 0.1015, 84),
     ("EMI-20L-9.15%-15y", 2000000, 0.0915, 180), ("EMI-zero-rate", 120000, 0, 12)], start=2):
    emi[f"A{i}"] = cid
    inp(emi, f"B{i}", p); inp(emi, f"C{i}", rate); inp(emi, f"D{i}", n)
    fx(emi, f"E{i}", f"=IF(C{i}=0,B{i}/D{i},PMT(C{i}/12,D{i},-B{i}))", expected=True).number_format = INR

# ----------------------------------------------------------------- Moratorium / subvention
mor = wb.create_sheet("Moratorium")
heads = ["id", "d1_amount", "d1_month", "d2_amount", "d2_month", "repay_month", "annual_rate", "subvention",
         "reduction", "tenure_months", "expected_capitalised", "expected_principal_at_repayment", "expected_emi"]
for col, h in enumerate(heads, start=1):
    mor.cell(row=1, column=col, value=h).font = BOLD
mor_cases = [
    ("MOR-capitalised", 250000, 0, 250000, 12, 36, 0.09, "none", 0, 60),
    ("MOR-PMV-3pct-subvention", 250000, 0, 250000, 12, 36, 0.09, "rate", 0.03, 60),
    ("MOR-PMUSP-full-subvention", 400000, 0, 400000, 12, 36, 0.0865, "full", 0, 120),
    ("MOR-SBI-student-1y", 300000, 0, 300000, 12, 48, 0.1015, "none", 0, 120),
]
for i, c in enumerate(mor_cases, start=2):
    mor[f"A{i}"] = c[0]
    for col, v in zip("BCDEFGHIJ", c[1:]):
        inp(mor, f"{col}{i}", v)
    eff = f'IF(H{i}="full",0,MAX(0,G{i}-IF(H{i}="rate",I{i},0)))'
    fx(mor, f"K{i}", f"=B{i}*{eff}*(F{i}-C{i})/12+D{i}*{eff}*(F{i}-E{i})/12", expected=True).number_format = INR
    fx(mor, f"L{i}", f"=B{i}+D{i}+K{i}", expected=True).number_format = INR
    fx(mor, f"M{i}", f"=PMT(G{i}/12,J{i},-L{i})", expected=True).number_format = INR

# ----------------------------------------------------------------- Section 129
s129 = wb.create_sheet("Sec129")
for col, h in enumerate(["id", "interest_paid", "parent_gross", "regime", "year_index", "max_years", "expected_saving"], start=1):
    s129.cell(row=1, column=col, value=h).font = BOLD
for i, (cid, interest, pg, regime, yi) in enumerate(
    [("S129-old-30pct", 100000, 1100000, "old", 1), ("S129-old-20pct", 80000, 800000, "old", 3),
     ("S129-new-zero", 100000, 1100000, "new", 1), ("S129-notsure-zero", 100000, 1100000, "not_sure", 1),
     ("S129-old-year9-zero", 100000, 1100000, "old", 9)], start=2):
    s129[f"A{i}"] = cid
    inp(s129, f"B{i}", interest); inp(s129, f"C{i}", pg); inp(s129, f"D{i}", regime)
    inp(s129, f"E{i}", yi); inp(s129, f"F{i}", 8)
    fx(s129, f"G{i}", f'=IF(AND(D{i}="old",E{i}>=1,E{i}<=F{i},B{i}>0),B{i}*{marginal_rate(f"C{i}", "old")},0)', expected=True).number_format = INR

# ----------------------------------------------------------------- FX
fxs = wb.create_sheet("FX")
for col, h in enumerate(["id", "inr_per_unit_base", "depreciation", "year", "amount_local", "expected_inr"], start=1):
    fxs.cell(row=1, column=col, value=h).font = BOLD
for i, (cid, base, dep, yr, amt) in enumerate(
    [("FX-EUR-y1", 95, 0.02, 1, 10000), ("FX-EUR-y5", 95, 0.02, 5, 10000), ("FX-USD-y8-4pct", 83, 0.04, 8, 60000)], start=2):
    fxs[f"A{i}"] = cid
    inp(fxs, f"B{i}", base); inp(fxs, f"C{i}", dep); inp(fxs, f"D{i}", yr); inp(fxs, f"E{i}", amt)
    fx(fxs, f"F{i}", f"=E{i}*B{i}*(1+C{i})^(D{i}-1)", expected=True).number_format = INR

# ----------------------------------------------------------------- ROC
roc = wb.create_sheet("ROC")
roc["A1"] = "rank"; roc["B1"] = "n"; roc["C1"] = "expected_weight"
for k in range(1, 7):
    inp(roc, f"A{k + 1}", k); inp(roc, f"B{k + 1}", 6)
    fx(roc, f"C{k + 1}", f"=SUMPRODUCT((ROW($A$1:$A$20)>=A{k + 1})*(ROW($A$1:$A$20)<=B{k + 1})/ROW($A$1:$A$20))/B{k + 1}", expected=True)

# ----------------------------------------------------------------- Plans (deterministic NPV at "likely")
PLAN_FIELDS = [
    ("start_age", 17), ("horizon_age", 35), ("gap_years", 0), ("gap_year_cost_inr", 0),
    ("stage_years", 4), ("stage_currency_fx", 1), ("tuition", 200000), ("living", 100000), ("one_time", 0),
    ("work_india", 1), ("work_fx", 1), ("start_gross_today", 900000), ("growth", 0.08), ("ai_delta", 0.005),
    ("effective_tax_rate", 0), ("living_after_today", 250000),
    ("switch_after_years", 99), ("ret_start_gross_today", 0), ("ret_growth", 0), ("ret_living_today", 0),
    ("infl_education", 0.08), ("infl_general", 0.04), ("infl_stage_local", 0.04), ("infl_work_local", 0.04),
    ("fx_depreciation", 0.02), ("discount_rate", 0.07),
    ("loan_amount", 0), ("loan_rate", 0), ("loan_tenure_years", 10), ("loan_moratorium_months", 12), ("loan_serviced", 0),
    ("parent_old_regime", 0), ("parent_gross", 1500000), ("s129_max_years", 8),
]

PLANS = {
    "NPV-core-india": {},
    "NPV-repeat-year": {"gap_years": 1, "gap_year_cost_inr": 150000},
    "NPV-fallback": {"start_gross_today": 450000, "growth": 0.06, "ai_delta": 0, "living_after_today": 200000},
    "NPV-abroad-return": {"stage_years": 2, "stage_currency_fx": 95, "tuition": 3000, "living": 11000, "one_time": 1000,
                          "work_india": 0, "work_fx": 95, "start_gross_today": 50000, "growth": 0.03, "ai_delta": 0,
                          "effective_tax_rate": 0.35, "living_after_today": 14000, "switch_after_years": 3,
                          "ret_start_gross_today": 1200000, "ret_growth": 0.07, "ret_living_today": 300000,
                          "infl_education": 0.03, "infl_stage_local": 0.02, "infl_work_local": 0.02},
    "NPV-custom-path": {"stage_years": 3, "tuition": 350000, "living": 180000, "one_time": 50000,
                        "start_gross_today": 600000, "growth": 0.07, "ai_delta": 0.01, "living_after_today": 240000},
    "NPV-loan-old-regime-s129": {"loan_amount": 800000, "loan_rate": 0.1015, "parent_old_regime": 1},
    "NPV-loan-serviced": {"loan_amount": 800000, "loan_rate": 0.0915, "loan_serviced": 1},
}

MONTHS = 400  # covers study + moratorium + tenure within the horizon


def build_plan(name, overrides):
    ws = wb.create_sheet(name[:31])
    ws["A1"] = name
    ws["A1"].font = BOLD
    cell = {}
    for i, (key, default) in enumerate(PLAN_FIELDS, start=2):
        ws[f"A{i}"] = key
        inp(ws, f"B{i}", overrides.get(key, default))
        cell[key] = f"$B${i}"
    c = cell
    # Derived
    ws["D2"] = "H"; fx(ws, "E2", f"={c['horizon_age']}-{c['start_age']}")
    ws["D3"] = "first_study_year"; fx(ws, "E3", f"={c['gap_years']}+1")
    ws["D4"] = "first_work_year"; fx(ws, "E4", f"={c['gap_years']}+{c['stage_years']}+1")
    ws["D5"] = "repay_start_month"; fx(ws, "E5", f"=(E4-1)*12+{c['loan_moratorium_months']}")
    ws["D6"] = "switch_year"; fx(ws, "E6", f"=E4+{c['switch_after_years']}")
    ws["D7"] = "expected_npv"
    H, FS, FW, RS, SW = "$E$2", "$E$3", "$E$4", "$E$5", "$E$6"

    # Yearly table from row 40
    top = 40
    heads = ["t", "study", "work", "fx_stage", "fx_work", "study_cost", "loan_disbursed", "out_of_pocket", "gap_cost",
             "switched", "gross", "tax", "in_hand", "living_after", "loan_payment", "loan_interest", "s129_index",
             "s129_saving", "cash_flow", "discounted"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=top - 1, column=j, value=h).font = BOLD
    col = {h: get_column_letter(j) for j, h in enumerate(heads, start=1)}
    for t in range(1, 19):
        r = top + t - 1
        g = lambda h: f"{col[h]}{r}"  # noqa: E731
        ws[g("t")] = t
        fx(ws, g("study"), f"=AND({g('t')}>={FS},{g('t')}<{FW})*1")
        fx(ws, g("work"), f"=({g('t')}>={FW})*({g('t')}<={H})*1")
        fx(ws, g("fx_stage"), f"=IF({c['stage_currency_fx']}=1,1,{c['stage_currency_fx']}*(1+{c['fx_depreciation']})^({g('t')}-1))")
        fx(ws, g("fx_work"), f"=IF({c['work_fx']}=1,1,{c['work_fx']}*(1+{c['fx_depreciation']})^({g('t')}-1))")
        fx(ws, g("study_cost"),
           f"={g('study')}*({c['tuition']}*(1+{c['infl_education']})^({g('t')}-1)+{c['living']}*(1+{c['infl_stage_local']})^({g('t')}-1)"
           f"+({g('t')}={FS})*{c['one_time']}*(1+{c['infl_stage_local']})^({g('t')}-1))*{g('fx_stage')}")
        prev_disb = f"SUM(${col['loan_disbursed']}${top}:{col['loan_disbursed']}{r - 1})" if t > 1 else "0"
        fx(ws, g("loan_disbursed"), f"={g('study')}*MIN({g('study_cost')},MAX(0,{c['loan_amount']}-{prev_disb}))")
        fx(ws, g("out_of_pocket"), f"={g('study_cost')}-{g('loan_disbursed')}")
        fx(ws, g("gap_cost"), f"=({g('t')}<{FS})*{c['gap_year_cost_inr']}*(1+{c['infl_general']})^({g('t')}-1)")
        fx(ws, g("switched"), f"={g('work')}*({g('t')}>={SW})*1")
        main_gross = (f"{c['start_gross_today']}*(1+{c['infl_work_local']})^({FW}-1)"
                      f"*(1+{c['growth']}-{c['ai_delta']})^({g('t')}-{FW})")
        ret_gross = (f"{c['ret_start_gross_today']}*(1+{c['infl_general']})^({SW}-1)"
                     f"*(1+{c['ret_growth']})^({g('t')}-{SW})")
        fx(ws, g("gross"), f"={g('work')}*IF({g('switched')}=1,{ret_gross},{main_gross})")
        india_now = f"OR({g('switched')}=1,{c['work_india']}=1)"
        fx(ws, g("tax"), f"=IF(AND({g('work')}=1,{india_now}),{income_tax(g('gross'), 'new')},0)")
        fx(ws, g("in_hand"),
           f"=IF({g('work')}=0,0,IF({india_now},{g('gross')}-{g('tax')},{g('gross')}*(1-{c['effective_tax_rate']})*{g('fx_work')}))")
        fx(ws, g("living_after"),
           f"={g('work')}*IF({g('switched')}=1,{c['ret_living_today']}*(1+{c['infl_general']})^({g('t')}-1),"
           f"{c['living_after_today']}*(1+{c['infl_work_local']})^({g('t')}-1)*IF({c['work_india']}=1,1,{g('fx_work')}))")
        loan_ws = f"'{name[:26]} loan'"
        fx(ws, g("loan_payment"), f"=SUMIF({loan_ws}!$B$2:$B${MONTHS + 1},{g('t')},{loan_ws}!$I$2:$I${MONTHS + 1})")
        fx(ws, g("loan_interest"), f"=SUMIF({loan_ws}!$B$2:$B${MONTHS + 1},{g('t')},{loan_ws}!$H$2:$H${MONTHS + 1})")
        prev_idx = f"{col['s129_index']}{r - 1}" if t > 1 else "0"
        fx(ws, g("s129_index"), f"=IF({prev_idx}>0,{prev_idx}+1,IF({g('loan_interest')}>0,1,0))")
        fx(ws, g("s129_saving"),
           f"=IF(AND({c['parent_old_regime']}=1,{g('s129_index')}>=1,{g('s129_index')}<={c['s129_max_years']},{g('loan_interest')}>0),"
           f"{g('loan_interest')}*{marginal_rate(c['parent_gross'], 'old')},0)")
        fx(ws, g("cash_flow"),
           f"={g('in_hand')}-{g('living_after')}-{g('out_of_pocket')}-{g('gap_cost')}-{g('loan_payment')}+{g('s129_saving')}")
        fx(ws, g("discounted"), f"={g('cash_flow')}/(1+{c['discount_rate']})^{g('t')}")
    fx(ws, "E7", f"=SUM({col['discounted']}{top}:{col['discounted']}{top + 17})", expected=True).number_format = INR

    # Monthly loan schedule (simple moratorium interest, capitalised unless serviced; annuity after)
    ls = wb.create_sheet(f"{name[:26]} loan")
    lheads = ["month", "plan_year", "disbursed", "balance_disbursed", "mor_interest", "capitalised_cum", "emi_balance_start",
              "interest_paid", "payment"]
    for j, h in enumerate(lheads, start=1):
        ls.cell(row=1, column=j, value=h).font = BOLD
    P = f"'{name[:31]}'!"
    rate = f"{P}{c['loan_rate'].replace('$', '')}"
    serviced = f"{P}{c['loan_serviced'].replace('$', '')}"
    repay = f"{P}E5"
    tenure_m = f"{P}{c['loan_tenure_years'].replace('$', '')}*12"
    # principal at repayment = total disbursed + capitalised (if not serviced)
    ls["K1"] = "principal_at_repayment"
    ls["K2"] = f"=SUM(C2:C{MONTHS + 1})+IF({serviced}=1,0,SUMIF(A2:A{MONTHS + 1},\"<\"&{repay},E2:E{MONTHS + 1}))"
    ls["L1"] = "emi"
    ls["L2"] = f"=IF(K2=0,0,PMT({rate}/12,{tenure_m},-K2))"
    for m in range(MONTHS):
        r = m + 2
        ls[f"A{r}"] = m
        ls[f"B{r}"] = f"=INT(A{r}/12)+1"
        # disbursed at the start of each study year (month = (t-1)*12) from the plan table
        ls[f"C{r}"] = (f"=IF(MOD(A{r},12)=0,IFERROR(INDEX({P}${col['loan_disbursed']}${top}:${col['loan_disbursed']}${top + 17},B{r}),0),0)")
        ls[f"D{r}"] = f"=SUM($C$2:C{r})"
        ls[f"E{r}"] = f"=IF(A{r}<{repay},D{r}*{rate}/12,0)"
        ls[f"F{r}"] = f"=SUM($E$2:E{r})"
        prev = f"G{r - 1}-(I{r - 1}-H{r - 1})" if m > 0 else "0"
        ls[f"G{r}"] = f"=IF(A{r}={repay},$K$2,IF(A{r}>{repay},{prev},0))"
        ls[f"H{r}"] = (f"=IF(A{r}<{repay},IF({serviced}=1,E{r},0),"
                       f"IF(A{r}<{repay}+{tenure_m},G{r}*{rate}/12,0))")
        ls[f"I{r}"] = (f"=IF(A{r}<{repay},IF({serviced}=1,E{r},0),"
                       f"IF(A{r}<{repay}+{tenure_m},$L$2,0))")


for name, ov in PLANS.items():
    build_plan(name, ov)

for ws in wb.worksheets:
    ws.column_dimensions["A"].width = 26

wb.save(sys.argv[1])
print(f"wrote {sys.argv[1]}")
