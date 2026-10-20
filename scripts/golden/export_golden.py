"""Export recalculated golden cases to JSON for the engine tests.

Usage: python3 -I scripts/golden/export_golden.py docs/golden/golden-cases.xlsx packages/engine/src/golden/golden-cases.json
Run after build_golden.py and LibreOffice recalculation. Never edit the JSON by hand.
"""
import hashlib
import json
import sys

from openpyxl import load_workbook

src, out = sys.argv[1], sys.argv[2]
wb = load_workbook(src, data_only=True)


def rows(name):
    ws = wb[name]
    heads = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        yield dict(zip(heads, r))


def need(v, where):
    if v is None:
        sys.exit(f"missing value at {where}: recalculate the workbook first")
    return v


rules_ws = wb["Rules"]
def slabs(regime):
    out_ = []
    for r in rules_ws.iter_rows(min_row=2, max_col=4, values_only=True):
        if r[0] == regime:
            out_.append({"from": r[1], "to": None if r[2] >= 1e14 else r[2], "rate": r[3]})
    return out_
param = {rules_ws[f"F{i}"].value: (rules_ws[f"G{i}"].value, rules_ws[f"H{i}"].value) for i in range(2, 6)}
rules = {
    regime: {
        "regime": regime,
        "slabs": slabs(regime),
        "standardDeduction": param["standard_deduction"][k],
        "rebate": {"maxTaxableIncome": param["rebate_max_taxable"][k], "maxRebate": param["rebate_max"][k]},
        "cessRate": param["cess"][k],
        "section129Allowed": regime == "old",
    }
    for k, regime in enumerate(["new", "old"])
}

cases = []
for r in rows("Tax"):
    cases.append({"id": r["id"], "category": "tax", "inputs": {"regime": r["regime"], "gross": r["gross"], "marginalRelief": bool(r["marginal_relief"])}, "expected": {"total": need(r["expected_total"], r["id"])}})
for r in rows("EMI"):
    cases.append({"id": r["id"], "category": "emi", "inputs": {"principal": r["principal"], "annualRate": r["annual_rate"], "months": r["months"]}, "expected": {"emi": need(r["expected_emi"], r["id"])}})
for r in rows("Moratorium"):
    cases.append({
        "id": r["id"], "category": "moratorium",
        "inputs": {k: r[k] for k in ["d1_amount", "d1_month", "d2_amount", "d2_month", "repay_month", "annual_rate", "subvention", "reduction", "tenure_months"]},
        "expected": {"capitalised": need(r["expected_capitalised"], r["id"]), "principalAtRepayment": r["expected_principal_at_repayment"], "emi": r["expected_emi"]},
    })
for r in rows("Sec129"):
    cases.append({"id": r["id"], "category": "section129", "inputs": {"interestPaid": r["interest_paid"], "parentGross": r["parent_gross"], "regime": r["regime"], "yearIndex": r["year_index"], "maxYears": r["max_years"]}, "expected": {"saving": need(r["expected_saving"], r["id"])}})
for r in rows("FX"):
    cases.append({"id": r["id"], "category": "fx", "inputs": {"base": r["inr_per_unit_base"], "depreciation": r["depreciation"], "year": r["year"], "amountLocal": r["amount_local"]}, "expected": {"inr": need(r["expected_inr"], r["id"])}})
roc = [r["expected_weight"] for r in rows("ROC")]
cases.append({"id": "ROC-6", "category": "roc", "inputs": {"n": 6}, "expected": {"weights": roc}})
for name in [n for n in wb.sheetnames if n.startswith("NPV-") and not n.endswith(" loan")]:
    ws = wb[name]
    inputs = {}
    i = 2
    while ws[f"A{i}"].value:
        inputs[ws[f"A{i}"].value] = ws[f"B{i}"].value
        i += 1
    cases.append({"id": name, "category": "npv", "inputs": inputs, "expected": {"npv": need(ws["E7"].value, name)}})

sha = hashlib.sha256(open(src, "rb").read()).hexdigest()
doc = {
    "status": "draft-awaiting-approval",
    "source": "docs/golden/golden-cases.xlsx",
    "sourceSha256": sha,
    "tolerance": {"absoluteInr": 1, "relative": 0.0001},
    "rules": rules,
    "cases": cases,
}
with open(out, "w") as f:
    json.dump(doc, f, indent=2)
    f.write("\n")
print(f"wrote {out}: {len(cases)} cases")
